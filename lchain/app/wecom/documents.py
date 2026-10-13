"""文件读取与解析：图片 / PDF / Word / 文本 -> deepseek-flash 解读 -> 存入 wecom_documents（“同步到 repo”）。

- 图片：直接以 image_url(base64) 发给多模态模型
- PDF：有文字层时抽取文本；扫描件（文字很少）把前 N 页渲染成图片交给模型识别
- docx / txt / md / csv：抽取文本后交给模型解读
"""
from __future__ import annotations

import asyncio
import base64
import logging
import mimetypes
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from langchain_core.runnables import RunnableConfig

from app.core.config import get_settings
from app.wecom import store
from app.wecom.llm import call_text

logger = logging.getLogger(__name__)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".log"}
MAX_TEXT_CHARS = 30000


@dataclass
class SavedFile:
    name: str
    path: str
    mime: str
    size: int

    @property
    def kind(self) -> str:
        return file_kind(self.name, self.mime)

    def as_attachment(self) -> dict[str, Any]:
        return {"name": self.name, "path": self.path, "mime": self.mime, "size": self.size, "kind": self.kind,
                "url": "/api/wecom/files/" + Path(self.path).name}


def file_kind(name: str, mime: str = "") -> str:
    ext = Path(name).suffix.lower()
    if ext in IMAGE_EXT or mime.startswith("image/"):
        return "image"
    if ext == ".pdf" or mime == "application/pdf":
        return "pdf"
    if ext in (".docx",):
        return "docx"
    if ext in TEXT_EXT or mime.startswith("text/"):
        return "text"
    return "other"


def upload_dir() -> Path:
    d = Path(get_settings().wecom_upload_dir)
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_bytes(name: str, data: bytes, mime: str | None = None) -> SavedFile:
    s = get_settings()
    if len(data) > s.wecom_max_file_mb * 1024 * 1024:
        raise ValueError(f"文件超过 {s.wecom_max_file_mb}MB")
    name = Path(name or "file").name
    ext = Path(name).suffix.lower()
    path = upload_dir() / f"{uuid.uuid4().hex}{ext}"
    path.write_bytes(data)
    mime = mime or mimetypes.guess_type(name)[0] or "application/octet-stream"
    return SavedFile(name=name, path=str(path), mime=mime, size=len(data))


def _b64_image(data: bytes, mime: str) -> dict[str, Any]:
    return {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{base64.b64encode(data).decode()}"}}


def _extract(att: dict[str, Any]) -> tuple[str, list[dict[str, Any]]]:
    """同步：返回 (抽取的文本, 图片 content blocks)。"""
    path, kind, mime = att["path"], att.get("kind") or file_kind(att["name"], att.get("mime", "")), att.get("mime", "")
    data = Path(path).read_bytes()
    if kind == "image":
        return "", [_b64_image(data, mime or "image/png")]
    if kind == "pdf":
        import pymupdf

        max_pages = get_settings().wecom_pdf_max_pages
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            pages = list(doc)[:max_pages]
            text = "\n\n".join(f"--- 第 {i + 1} 页 ---\n{p.get_text().strip()}" for i, p in enumerate(pages))
            total = doc.page_count
            # 平均每页文字很少 -> 视为扫描件，渲染成图片
            if len(text.replace("-", "").strip()) < 50 * max(1, len(pages)):
                imgs = [_b64_image(p.get_pixmap(dpi=110).tobytes("jpeg"), "image/jpeg") for p in pages]
                return f"（扫描版 PDF，共 {total} 页，以下为前 {len(pages)} 页图片）", imgs
        if total > max_pages:
            text += f"\n\n（共 {total} 页，仅解析前 {max_pages} 页）"
        return text[:MAX_TEXT_CHARS], []
    if kind == "docx":
        import docx

        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)[:MAX_TEXT_CHARS], []
    if kind == "text":
        return data.decode("utf-8", errors="ignore")[:MAX_TEXT_CHARS], []
    raise ValueError(f"暂不支持的文件类型：{att['name']}")


async def parse_attachment(att: dict[str, Any], *, group_id: str, message_id: int, user_text: str = "",
                           cfg: RunnableConfig | None = None, overrides: dict[str, str] | None = None) -> dict[str, Any]:
    """解析单个附件并落库，返回文档记录（dict）。失败不抛异常，status=failed。"""
    raw_text, parsed, status, error = "", "", "ok", None
    try:
        raw_text, images = await asyncio.to_thread(_extract, att)
        head = f"文件名：{att['name']}\n用户随文件发送的话：{user_text or '（无）'}\n"
        if images:
            content: str | list[dict[str, Any]] = [{"type": "text", "text": head + (raw_text or "请解读这份文件。")}, *images]
        else:
            if not raw_text.strip():
                raise ValueError("未能从文件中提取到文字内容")
            content = head + "\n文件内容：\n" + raw_text
        parsed, _ = await call_text("wecom-doc-parse", content, cfg=cfg, overrides=overrides)
    except ValueError as e:
        status, error = ("unsupported" if "不支持" in str(e) else "failed"), str(e)
    except Exception as e:  # noqa: BLE001
        logger.exception("文件解析失败: %s", att.get("name"))
        status, error = "failed", f"{type(e).__name__}: {e}"
    d = await store.add_document(group_id=group_id, message_id=message_id, file_name=att["name"],
                                 mime_type=att.get("mime", ""), file_path=att["path"], raw_text=raw_text,
                                 parsed_text=parsed, status=status, error=error)
    return {"id": d.id, "file_name": d.file_name, "status": status, "error": error, "parsed_text": parsed}
